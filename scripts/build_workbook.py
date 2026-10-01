"""Build the monochrome Meiryo UI workbook from the complete Markdown guide."""
from __future__ import annotations
from copy import copy
import json
import math
from pathlib import Path
import re
import unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
FONT = "Meiryo UI"
DARK = "242424"
MID = "E5E5E5"
LIGHT = "F5F5F5"
WHITE = "FFFFFF"
TEXT = "222222"
LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")


def width(text):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(text))


def lines(text, columns):
    return sum(max(1, math.ceil(width(line) / columns)) for line in str(text).split("\n"))


def cell_style(cell, bold=False, fill=None, size=11):
    cell.font = Font(name=FONT, size=size, color=TEXT, bold=bold)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = Border(bottom=Side(style="hair", color="D0D0D0"))
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if isinstance(cell.value, str) and cell.data_type == "f":
        cell.data_type = "s"


def setup(ws, title):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "666666"
    for column, size in (("A", 7), ("B", 27), ("C", 76), ("D", 57)):
        ws.column_dimensions[column].width = size
    ws.merge_cells("A1:D1")
    ws["A1"] = title
    cell_style(ws["A1"], True, DARK, 17)
    ws["A1"].font = Font(name=FONT, size=17, color=WHITE, bold=True)
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:D2")
    ws["A2"] = "確認日 2026-10-01  |  JMeter 5.6.3  |  Cognito M2M  |  全文は同梱Markdownと共通"
    cell_style(ws["A2"], False, LIGHT, 10)
    ws.row_dimensions[2].height = 25
    ws.freeze_panes = "B4"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.oddFooter.center.text = '&"Meiryo UI,Regular"&8 &A   &P / &N'


def main():
    markdown = (DOCS / "Cognito_JMeter_詳細ガイド.md").read_text(encoding="utf-8")
    sources = []
    source_index = {}
    for match in LINK_RE.finditer(markdown):
        title, url = match.groups()
        if url not in source_index:
            source_index[url] = len(sources) + 1
            sources.append({"id": f"S{len(sources)+1:02}", "title": title, "url": url, "checked": "2026-10-01"})
    (DOCS / "sources.json").write_text(json.dumps(sources, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def plain(text):
        return LINK_RE.sub(lambda m: m[1] + f" [S{source_index[m[2]]:02}]", text).replace("**", "").replace(chr(96), "")

    wb = Workbook()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    wb._named_styles[0].font = Font(name=FONT, size=11)
    cover = wb.active
    cover.title = "00_案内"
    setup(cover, "Cognito M2M × API Gateway × JMeter 5.6.3")
    cover.append(["No.", "内容", "説明", "リンク・備考"])
    for cell in cover[3]:
        cell_style(cell, True, MID)
    info = [
        ("対象", "専用M2Mクライアント、トークンの再利用、OpenAPIからのJMX生成", "2026-10-01公式情報確認"),
        ("資料の読み方", "01〜11で原理、12〜20で実装と操作、21〜24で境界と検証を確認。全章の本文・表・コードを収録。", "Markdownと共通内容"),
        ("初期値", "API localhost:8080、1スレッド、1ループ、有効API1件。Cognito取得は追加のHTTP1件。", "X-API-KEY=XXXXXXXXXXXX"),
        ("実装", "Bash / PowerShellの入口 + 共通Python。JMXは単独でJMeterに読込可能。", "Python 3.11+が生成に必要"),
        ("最新仕様", "2025-11 M2Mクライアント料金項目廃止、2026-02 secret rotation、2026-08 GetClientToken。", "本実装は/oauth2/token"),
        ("検証の境界", "実AWS・実OpenAPIは未提供。ローカル模擬サーバーでJMeter実行。実施結果はVALIDATION.md。", "RHEL9.8実機は未検証"),
        ("フォント", "全使用セルをMeiryo UI。フォント未導入機では代替表示になる可能性。", "フォント本体は非同梱"),
    ]
    for index, (key, body, note) in enumerate(info, 1):
        cover.append([index, key, body, note])
        for c in cover[cover.max_row]:
            cell_style(c, fill=LIGHT if index % 2 else WHITE)
        cover.row_dimensions[cover.max_row].height = 60
    cover.append(["No.", "章別索引", "クリックすると該当シートへ移動", "リンク"])
    for c in cover[cover.max_row]:
        cell_style(c, True, MID)
    cover.row_dimensions[cover.max_row].height = 30
    sections = re.split(r"(?m)^## (\d{2} .+)\n", markdown)
    for title, body in zip(sections[1::2], sections[2::2]):
        sheet_title = re.sub(r"[\[\]:*?/\\]", "", title.replace(" ", "_", 1))[:31]
        ws = wb.create_sheet(sheet_title)
        setup(ws, title)
        row = 3
        block = False
        seq = 0
        table_headers = None
        paragraphs = []

        def text_row(value, category="本文", fill=None, bold=False):
            nonlocal row, seq
            seq += 1
            ws.cell(row, 1, seq)
            # Prose/code uses the full content width. Tables keep real separate columns.
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            ws.cell(row, 2, value)
            for cells in ws.iter_rows(min_row=row, max_row=row, max_col=4):
                for c in cells:
                    cell_style(c, bold, fill, 10 if category == "コード" else 11)
            height = max(23, (lines(value, 145) + 1) * (14 if category == "コード" else 16))
            ws.row_dimensions[row].height = min(390, height)
            row += 1

        def flush():
            if paragraphs:
                value = plain(" ".join(paragraphs))
                # Short chunks prevent Excel's 409-point row-height ceiling from clipping.
                while value:
                    text_row(value[:550])
                    value = value[550:]
                paragraphs.clear()

        for line in body.splitlines():
            if line.startswith("~~~"):
                flush()
                block = not block
                if block:
                    text_row("コード / " + line[3:], "コード", MID, True)
                continue
            if block:
                text_row(line or " ", "コード", LIGHT)
                continue
            if line.startswith("|"):
                flush()
                parts = [plain(x.strip()) for x in line.strip().strip("|").split("|")]
                if all(re.fullmatch(r"[- :]+", x or "-") for x in parts):
                    continue
                is_header = table_headers is None
                if is_header:
                    table_headers = parts
                values = [seq + 1] + parts[:3]
                while len(values) < 4:
                    values.append("")
                if len(parts) > 3:
                    values[3] += "\n" + "\n".join(parts[3:])
                seq += 1
                for col, value in enumerate(values, 1):
                    c = ws.cell(row, col, value)
                    cell_style(c, is_header, MID if is_header else LIGHT if seq % 2 else WHITE)
                row_lines = max(lines(values[1], 24), lines(values[2], 70), lines(values[3], 51))
                ws.row_dimensions[row].height = max(29, min(390, (row_lines + 1) * 16))
                row += 1
                continue
            if not line.strip():
                flush()
                table_headers = None
                continue
            if re.match(r"^\d+\.\s", line):
                flush()
                text_row(plain(line))
            else:
                paragraphs.append(line)
        flush()
        ws.print_area = f"A1:D{max(3,row-1)}"
        cover.append([title[:2], sheet_title, title[3:], "シートを開く"])
        for c in cover[cover.max_row]:
            cell_style(c, fill=LIGHT if cover.max_row % 2 else WHITE)
        cover.cell(cover.max_row, 4).hyperlink = f"#'{sheet_title}'!A1"
        cover.cell(cover.max_row, 4).font = Font(name=FONT, size=11, color="333333", underline="single")
        cover.row_dimensions[cover.max_row].height = 32

    calc = wb.create_sheet("25_発行数試算")
    setup(calc, "発行数の概算（JMeter設定を変更する機能ではありません）")
    rows = [
        ["区分", "入力・計算項目", "値", "意味 / 前提"],
        ["入力", "キャッシュ数 C", 100, "本実装では総スレッド数。エンジン合計で考える"],
        ["入力", "トークン寿命 L 秒", 3600, "Cognito応答のexpires_in"],
        ["入力", "再利用率 r", 0.75, "実装の初期値。jitterをこの計算では無視"],
        ["入力", "試験時間 D 秒", 3600, "全スレッド同時開始・定常的に呼ぶ仮定"],
        ["入力", "全体API RPS", 100, "スレッド数ではなくAPI全体の送信頻度"],
        ["計算", "APIリクエスト数", "=C7*C8", "RPS × 試験秒数"],
        ["計算", "毎回取得時の発行数", "=C9", "APIごとに1回取得する仮定"],
        ["計算", "再利用時の発行数", "=C4*ROUNDUP(C7/(C5*C6),0)", "初回を含む概算。API呼出が十分頻繁な場合"],
        ["計算", "発行数の削減率", "=1-C11/C10", "毎回取得と比べた目安"],
        ["計算", "定常時の発行RPS", "=C4/(C5*C6)", "初回バーストは別。上限保証ではない"],
        ["条件", "この計算の境界", "jitter・ランプ・失敗・停止・疎な呼出を除外", "実測値や請求の保証ではありません"],
    ]
    for row in rows:
        calc.append(row)
        for c in calc[calc.max_row]:
            cell_style(c, calc.max_row == 3, MID if calc.max_row == 3 else LIGHT)
        calc.row_dimensions[calc.max_row].height = 44
    # Restore intentionally authored calculator formulas after text-cell styling.
    for index, row in enumerate(rows, 3):
        if isinstance(row[2], str) and row[2].startswith("="):
            calc.cell(index, 3, row[2])
            calc.cell(index, 3).data_type = "f"
    for r in (6, 12):
        calc.cell(r, 3).number_format = "0.0%"
    calc["C13"].number_format = "0.000"
    positive = DataValidation(type="decimal", operator="greaterThan", formula1="0")
    positive.errorTitle = "正の数を入力"
    positive.error = "0より大きい数値を指定してください。"
    positive.showErrorMessage = True
    calc.add_data_validation(positive)
    for address in ("C5", "C7", "C8"):
        positive.add(address)
    whole = DataValidation(type="whole", operator="greaterThan", formula1="0")
    whole.showErrorMessage = True
    calc.add_data_validation(whole)
    whole.add("C4")
    fraction = DataValidation(type="decimal", operator="between", formula1="0.01", formula2="0.99")
    fraction.showErrorMessage = True
    calc.add_data_validation(fraction)
    fraction.add("C6")
    calc.print_area = "A1:D14"

    refs = wb.create_sheet("26_公式出典")
    setup(refs, "公式出典と確認日")
    refs.append(["ID", "文書", "URL", "確認日"])
    for c in refs[3]:
        cell_style(c, True, MID)
    for item in sources:
        refs.append([item["id"], item["title"], item["url"], item["checked"]])
        row = refs.max_row
        for c in refs[row]:
            cell_style(c, fill=LIGHT if row % 2 else WHITE)
        refs.cell(row, 3).hyperlink = item["url"]
        refs.cell(row, 3).font = Font(name=FONT, size=10, color="333333", underline="single")
        refs.row_dimensions[row].height = max(48, (lines(item["url"], 65) + 1) * 15)
    refs.auto_filter.ref = f"A3:D{refs.max_row}"
    refs.print_area = f"A1:D{refs.max_row}"
    cover.print_area = f"A1:D{cover.max_row}"
    cover.auto_filter.ref = f"A11:D{cover.max_row}"
    # Apply font even to used blank/merged cells for consistent formatting metadata.
    for ws in wb:
        for row in ws:
            for c in row:
                font = copy(c.font)
                font.name = FONT
                c.font = font
                c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.sheet_view.zoomScale = 80
    output = DOCS / "Cognito_JMeter_詳細ガイド.xlsx"
    wb.save(output)
    check = load_workbook(output)
    assert len(check.worksheets) == 27
    assert all(c.font.name == FONT for ws in check for row in ws for c in row if c.value is not None)
    assert all(c.alignment.wrap_text for ws in check for row in ws for c in row if c.value is not None)
    assert sum(c.data_type == "f" for ws in check for row in ws for c in row) == 5
    report = {"sheets": len(check.worksheets), "sources": len(sources),
              "nonempty_cells": sum(c.value is not None for ws in check for row in ws for c in row),
              "font": FONT, "all_nonempty_cells_font_checked": True, "wrap_text_checked": True,
              "calculator_formulas": 5, "markdown_characters": len(markdown)}
    (DOCS / "workbook-validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
